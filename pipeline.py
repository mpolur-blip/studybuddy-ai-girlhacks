"""
StudyBuddy AI - text extraction, cleanup, chunking, and FAISS retrieval.
GirlHacks 2026 - built solo, from scratch, at the event.
"""

import re
import uuid
from dataclasses import dataclass, field
from typing import Optional

import numpy as np


# ---------------------------------------------------------------------------
# 1. Data model - one chunk = one retrievable piece of a document
# ---------------------------------------------------------------------------

@dataclass
class Chunk:
    chunk_id: str
    text: str
    source_doc: str
    location: str  # e.g. "slide 14" or "page 3" - lets the UI cite its source
    embedding: Optional[np.ndarray] = field(default=None, repr=False)


# ---------------------------------------------------------------------------
# 2. Extraction - pull raw text out of PDFs and slide decks
# ---------------------------------------------------------------------------

def extract_pdf(path):
    """Returns a list of (label, text) tuples, one per page."""
    import pdfplumber

    pages = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""  # scanned/image pages can return None
            pages.append((f"page {i}", text))
    return pages


def extract_pptx(path):
    """Returns a list of (label, text) tuples, one per slide."""
    from pptx import Presentation

    slides = []
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        slides.append((f"slide {i}", "\n".join(parts)))
    return slides


# ---------------------------------------------------------------------------
# 3. Cleanup
# ---------------------------------------------------------------------------

def clean_text(raw):
    """Strip blank-line clutter and collapse repeated whitespace."""
    text = re.sub(r"[ \t]+", " ", raw)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------------------
# 4. Chunking - slides are natural chunks; PDFs get a fixed-size fallback
# ---------------------------------------------------------------------------

def sliding_window(text, max_words=200, overlap_ratio=0.2):
    """Fixed-size word window with overlap - the safe fallback for plain prose."""
    words = text.split()
    if not words:
        return []
    step = max(1, int(max_words * (1 - overlap_ratio)))
    chunks = []
    for start in range(0, len(words), step):
        window = words[start:start + max_words]
        if window:
            chunks.append(" ".join(window))
        if start + max_words >= len(words):
            break
    return chunks


def chunk_document(pages, source_doc, is_slides):
    """
    Slide decks: one slide = one chunk (a slide is already one idea).
    PDFs: sliding window per page, since we're skipping heading-detection
    for time - simpler and reliable beats clever under a deadline.
    """
    chunks = []
    for label, text in pages:
        cleaned = clean_text(text)
        if not cleaned:
            continue

        pieces = [cleaned] if is_slides else sliding_window(cleaned)
        for piece in pieces:
            chunks.append(Chunk(
                chunk_id=str(uuid.uuid4())[:8],
                text=piece,
                source_doc=source_doc,
                location=label,
            ))
    return chunks


# ---------------------------------------------------------------------------
# 5. Embedding + FAISS retrieval
# ---------------------------------------------------------------------------

class ChunkIndex:
    """Wraps a Sentence-Transformers model + a FAISS index over a list of Chunks."""

    def __init__(self, model_name="all-MiniLM-L6-v2"):
        from sentence_transformers import SentenceTransformer

        self.model = SentenceTransformer(model_name)
        self.index = None
        self.chunks = []

    def build(self, chunks):
        import faiss

        if not chunks:
            raise ValueError("No chunks to index - check extraction/chunking upstream.")

        texts = [c.text for c in chunks]
        embeddings = self.model.encode(
            texts, convert_to_numpy=True, normalize_embeddings=True
        ).astype("float32")

        dim = embeddings.shape[1]
        self.index = faiss.IndexFlatIP(dim)  # normalized vectors -> inner product = cosine similarity
        self.index.add(embeddings)

        for chunk, vec in zip(chunks, embeddings):
            chunk.embedding = vec
        self.chunks = chunks

    def search(self, query, top_k=4):
        if self.index is None:
            raise RuntimeError("Call build() before search().")

        query_vec = self.model.encode(
            [query], convert_to_numpy=True, normalize_embeddings=True
        ).astype("float32")

        scores, indices = self.index.search(query_vec, top_k)
        results = []
        for score, idx in zip(scores[0], indices[0]):
            if idx == -1:
                continue
            results.append((self.chunks[idx], float(score)))
        return results


if __name__ == "__main__":
    # Quick manual test - run `python pipeline.py yourfile.pdf` to sanity-check
    import sys
    if len(sys.argv) > 1:
        path = sys.argv[1]
        is_slides = path.lower().endswith(".pptx")
        pages = extract_pptx(path) if is_slides else extract_pdf(path)
        chunks = chunk_document(pages, source_doc=path, is_slides=is_slides)
        print(f"Extracted {len(chunks)} chunks")

        index = ChunkIndex()
        index.build(chunks)
        print("Index built. Try a search:")
        hits = index.search("test query", top_k=2)
        for chunk, score in hits:
            print(f"  [{chunk.location}] score={score:.2f}: {chunk.text[:80]}")
